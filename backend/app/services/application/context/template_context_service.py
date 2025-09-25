"""
Template Context Service

模板上下文构建器，专注于占位符与段落文本匹配，为Agent提供精确的语境信息
核心功能：将解析出的占位符匹配到所在段落，提供业务语境上下文
"""

from __future__ import annotations

import re
import logging
from typing import Dict, Any, List, Optional
from dataclasses import dataclass
from datetime import datetime

logger = logging.getLogger(__name__)


@dataclass
class PlaceholderContext:
    """占位符上下文信息"""
    name: str
    placeholder_type: str  # "统计类", "周期类", "图表类"
    context_paragraph: str  # 占位符所在的完整自然段
    surrounding_text: str = ""  # 周围相关文本片段
    position_info: Dict[str, Any] = None  # 位置信息（段落索引、句子索引等）
    semantic_type: Optional[str] = None   # 细分语义类型：ranking/compare/period/stat/chart
    parsed_params: Optional[Dict[str, Any]] = None  # 解析参数（如top_n）

    def __post_init__(self):
        if self.position_info is None:
            self.position_info = {}
        if self.parsed_params is None:
            self.parsed_params = {}

    def to_agent_format(self) -> Dict[str, Any]:
        """转换为Agent友好的格式"""
        return {
            "placeholder_name": self.name,
            "type": self.placeholder_type,
            "context_paragraph": self.context_paragraph,
            "surrounding_text": self.surrounding_text,
            "position_info": self.position_info,
            "context_length": len(self.context_paragraph),
            "semantic_type": self.semantic_type,
            "parsed_params": self.parsed_params
        }


@dataclass
class TemplateContextInfo:
    """模板上下文信息 - 专注于语境提供"""
    template_id: str
    name: str
    description: str

    # 核心：占位符上下文列表
    placeholder_contexts: List[PlaceholderContext] = None

    # 模板元数据
    total_placeholders: int = 0
    parsed_successfully: bool = False

    # 审计信息
    created_at: Optional[datetime] = None
    updated_at: Optional[datetime] = None

    def __post_init__(self):
        if self.placeholder_contexts is None:
            self.placeholder_contexts = []

        self.total_placeholders = len(self.placeholder_contexts)

    def to_agent_format(self) -> Dict[str, Any]:
        """转换为Agent友好的格式"""
        return {
            "template_info": {
                "id": self.template_id,
                "name": self.name,
                "description": self.description
            },
            "placeholder_contexts": [pc.to_agent_format() for pc in self.placeholder_contexts],
            "summary": {
                "total_placeholders": self.total_placeholders,
                "统计类_count": len([pc for pc in self.placeholder_contexts if pc.placeholder_type == "统计类"]),
                "周期类_count": len([pc for pc in self.placeholder_contexts if pc.placeholder_type == "周期类"]),
                "图表类_count": len([pc for pc in self.placeholder_contexts if pc.placeholder_type == "图表类"]),
                "parsed_successfully": self.parsed_successfully
            }
        }


class TemplateContextBuilder:
    """模板上下文构建器 - 与实际数据库集成"""

    def __init__(self, container=None):
        self.container = container
        # 尝试获取数据库session
        self._db_session = None
        if container and hasattr(container, 'get_db'):
            try:
                self._db_session = next(container.get_db())
            except Exception as e:
                logger.warning(f"Could not get database session from container: {e}")

        # 导入实际的模型
        try:
            from app.models.template import Template
            from app.models.template_placeholder import TemplatePlaceholder
            self.Template = Template
            self.TemplatePlaceholder = TemplatePlaceholder
        except ImportError as e:
            logger.error(f"Could not import template models: {e}")
            self.Template = None
            self.TemplatePlaceholder = None

    async def build_template_context(
        self,
        user_id: str,
        template_id: str,
        force_refresh: bool = False
    ) -> Dict[str, Any]:
        """
        基于真实数据构建模板上下文信息（不使用mock/predefined数据）

        Args:
            user_id: 用户ID
            template_id: 模板ID
            force_refresh: 是否强制刷新缓存

        Returns:
            包含占位符语境信息的字典
        """
        try:
            logger.info(f"Building template context for {template_id}")

            # 1) 必须基于真实数据（数据库或文档），否则返回失败
            if not self.Template:
                return {"success": False, "error": "template_models_unavailable"}

            # 获取数据库会话
            db = self._db_session
            if not db and self.container and hasattr(self.container, 'get_db'):
                try:
                    db = next(self.container.get_db())
                except Exception as e:
                    logger.warning(f"Failed to acquire DB session: {e}")

            if not db:
                return {"success": False, "error": "db_session_unavailable"}

            # 查询模板与占位符
            from sqlalchemy.orm import Session
            from app.crud.crud_template import crud_template
            from app.crud.crud_template_placeholder import template_placeholder as crud_tpl_ph

            if not isinstance(db, Session):
                return {"success": False, "error": "invalid_db_session"}

            tpl = crud_template.get(db, id=template_id)
            if not tpl:
                return {"success": False, "error": "template_not_found"}

            placeholders = crud_tpl_ph.get_by_template(db, template_id=template_id, include_inactive=False)

            # 2) 生成占位符上下文（若无法解析文档，仅提供名称与类型，段落留空）
            placeholder_contexts: List[PlaceholderContext] = []
            doc_content: Optional[Dict[str, Any]] = None

            # 可选：尝试读取文档并做段落匹配（需要python-docx，文件路径需可访问）
            try:
                if tpl.file_path and hasattr(self, 'storage') and getattr(self, 'storage') is not None:
                    doc_content = await self._extract_document_content(tpl.file_path)
            except Exception as e:
                # 降低日志级别，避免无意义的告警
                logger.info(f"Document content extraction not available: {e}")
                doc_content = None

            for ph in placeholders:
                ph_name = getattr(ph, 'placeholder_name', None) or getattr(ph, 'placeholder_text', '')
                ph_type = self._classify_placeholder_type(ph_name or '')
                semantic_type = self._infer_semantic_type(ph_name or '')
                params = {}
                top_n = self._extract_top_n(ph_name or '')
                if top_n:
                    params['top_n'] = top_n

                paragraph_ctx = {"paragraph": "", "surrounding": "", "position": {"status": "unparsed"}}
                if doc_content and ph_name:
                    try:
                        paragraph_ctx = self._find_placeholder_context(ph_name, doc_content)
                    except Exception:
                        paragraph_ctx = {"paragraph": "", "surrounding": "", "position": {"status": "not_found"}}

                placeholder_contexts.append(
                    PlaceholderContext(
                        name=ph_name or "",
                        placeholder_type=ph_type,
                        context_paragraph=paragraph_ctx.get("paragraph", ""),
                        surrounding_text=paragraph_ctx.get("surrounding", ""),
                        position_info=paragraph_ctx.get("position", {"status": "unparsed"}),
                        semantic_type=semantic_type,
                        parsed_params=params
                    )
                )

            # 3) 组装上下文
            context_info = TemplateContextInfo(
                template_id=str(tpl.id),
                name=tpl.name or "",
                description=tpl.description or "",
                placeholder_contexts=placeholder_contexts,
                parsed_successfully=doc_content is not None,
                created_at=tpl.created_at or datetime.now(),
                updated_at=tpl.updated_at or datetime.now()
            )

            return {"success": True, "template_context": context_info.to_agent_format()}

        except Exception as e:
            logger.error(f"Failed to build template context for {template_id}: {e}")
            return {"success": False, "error": f"Template context building failed: {str(e)}"}

    async def _parse_and_match_context(
        self,
        file_path: str,
        context_info: TemplateContextInfo
    ) -> bool:
        """
        解析模板并匹配占位符语境

        Args:
            file_path: 模板文件路径
            context_info: 要填充的上下文信息对象

        Returns:
            是否成功解析
        """
        try:
            # 调用模板解析服务
            parse_result = await self.template_parser.parse(self.storage, file_path)

            if not parse_result.get("success"):
                logger.warning(f"Template parsing failed: {parse_result.get('error', 'Unknown error')}")
                return False

            # 获取解析出的占位符列表
            variables = parse_result.get("variables", [])

            # 获取原始文档内容用于上下文匹配
            doc_content = await self._extract_document_content(file_path)

            if not doc_content:
                logger.warning("Could not extract document content for context matching")
                # 即使没有内容，也创建基本的占位符记录
                for var_name in variables:
                    placeholder_context = PlaceholderContext(
                        name=var_name,
                        placeholder_type=self._classify_placeholder_type(var_name),
                        context_paragraph=f"占位符 {var_name} - 无可用上下文信息",
                        surrounding_text="",
                        position_info={"status": "no_content_available"}
                    )
                    context_info.placeholder_contexts.append(placeholder_context)

                context_info.parsed_successfully = len(variables) > 0
                return True

            # 为每个占位符匹配上下文段落
            for var_name in variables:
                paragraph_context = self._find_placeholder_context(var_name, doc_content)

                placeholder_context = PlaceholderContext(
                    name=var_name,
                    placeholder_type=self._classify_placeholder_type(var_name),
                    context_paragraph=paragraph_context.get("paragraph", ""),
                    surrounding_text=paragraph_context.get("surrounding", ""),
                    position_info=paragraph_context.get("position", {})
                )

                context_info.placeholder_contexts.append(placeholder_context)

            context_info.parsed_successfully = True
            return True

        except Exception as e:
            logger.error(f"Error in _parse_and_match_context: {e}")
            return False

    async def _extract_document_content(self, file_path: str) -> Optional[Dict[str, Any]]:
        """
        提取文档内容用于上下文分析

        Returns:
            包含段落和文本信息的字典
        """
        try:
            # 读取文档内容
            doc_bytes = await self.storage.read_bytes(file_path)
            if not doc_bytes:
                return None

            # 使用python-docx提取文档结构化内容
            from docx import Document
            from io import BytesIO

            doc = Document(BytesIO(doc_bytes))

            paragraphs = []
            for i, paragraph in enumerate(doc.paragraphs):
                if paragraph.text.strip():  # 只保留有内容的段落
                    paragraphs.append({
                        "index": i,
                        "text": paragraph.text.strip(),
                        "length": len(paragraph.text.strip())
                    })

            return {
                "paragraphs": paragraphs,
                "total_paragraphs": len(paragraphs)
            }

        except Exception as e:
            logger.warning(f"Failed to extract document content: {e}")
            return None

    def _find_placeholder_context(
        self,
        placeholder_name: str,
        doc_content: Dict[str, Any]
    ) -> Dict[str, Any]:
        """
        找到占位符所在的上下文段落

        Args:
            placeholder_name: 占位符名称（如"退货金额"）
            doc_content: 文档内容字典

        Returns:
            上下文信息字典
        """
        if not doc_content or not doc_content.get("paragraphs"):
            return {
                "paragraph": f"未找到占位符 {placeholder_name} 的上下文信息",
                "surrounding": "",
                "position": {"status": "not_found"}
            }

        paragraphs = doc_content["paragraphs"]

        # 查找包含占位符的段落
        # 占位符在Word中通常以{{ placeholder_name }}的形式出现
        placeholder_patterns = [
            f"{{{{{placeholder_name}}}}}",  # {{退货金额}}
            f"{{{placeholder_name}}}",     # {退货金额}
            placeholder_name               # 退货金额 (直接匹配)
        ]

        for para_info in paragraphs:
            paragraph_text = para_info["text"]

            # 检查是否包含占位符
            for pattern in placeholder_patterns:
                if pattern in paragraph_text:
                    # 找到了包含占位符的段落
                    surrounding_paras = self._get_surrounding_paragraphs(
                        para_info["index"], paragraphs
                    )

                    return {
                        "paragraph": paragraph_text,
                        "surrounding": surrounding_paras,
                        "position": {
                            "paragraph_index": para_info["index"],
                            "paragraph_length": para_info["length"],
                            "match_pattern": pattern,
                            "status": "found"
                        }
                    }

        # 如果没有直接找到，尝试语义相关匹配
        semantic_match = self._find_semantic_context(placeholder_name, paragraphs)
        if semantic_match:
            return semantic_match

        # 都没找到，返回默认信息
        return {
            "paragraph": f"未在文档中找到占位符 {placeholder_name} 的具体上下文，可能需要检查占位符格式",
            "surrounding": "",
            "position": {"status": "not_found_in_document"}
        }

    def _get_surrounding_paragraphs(
        self,
        target_index: int,
        paragraphs: List[Dict[str, Any]],
        window_size: int = 1
    ) -> str:
        """
        获取目标段落周围的相关段落文本

        Args:
            target_index: 目标段落索引
            paragraphs: 所有段落列表
            window_size: 前后获取的段落数量

        Returns:
            周围段落的文本
        """
        start_idx = max(0, target_index - window_size)
        end_idx = min(len(paragraphs), target_index + window_size + 1)

        surrounding_texts = []
        for i in range(start_idx, end_idx):
            if i != target_index and i < len(paragraphs):
                surrounding_texts.append(paragraphs[i]["text"])

        return " ".join(surrounding_texts)

    def _find_semantic_context(
        self,
        placeholder_name: str,
        paragraphs: List[Dict[str, Any]]
    ) -> Optional[Dict[str, Any]]:
        """
        基于语义相关性查找上下文

        Args:
            placeholder_name: 占位符名称
            paragraphs: 段落列表

        Returns:
            语义匹配的上下文信息
        """
        # 简单的关键词匹配策略
        keywords = self._extract_keywords_from_placeholder(placeholder_name)

        best_match = None
        best_score = 0

        for para_info in paragraphs:
            paragraph_text = para_info["text"]
            score = 0

            # 计算关键词匹配分数
            for keyword in keywords:
                if keyword in paragraph_text:
                    score += 1

            if score > best_score:
                best_score = score
                best_match = para_info

        if best_match and best_score > 0:
            surrounding_paras = self._get_surrounding_paragraphs(
                best_match["index"], paragraphs
            )

            return {
                "paragraph": best_match["text"],
                "surrounding": surrounding_paras,
                "position": {
                    "paragraph_index": best_match["index"],
                    "paragraph_length": best_match["length"],
                    "match_type": "semantic",
                    "match_score": best_score,
                    "status": "semantic_match"
                }
            }

        return None

    def _extract_keywords_from_placeholder(self, placeholder_name: str) -> List[str]:
        """
        从占位符名称中提取关键词

        Args:
            placeholder_name: 占位符名称

        Returns:
            关键词列表
        """
        # 基本的关键词提取
        keywords = [placeholder_name]

        # 如果是中文，尝试分词或者提取部分词汇
        if any('\u4e00' <= char <= '\u9fff' for char in placeholder_name):
            # 中文占位符，提取可能的关键词
            common_words = ["金额", "数量", "比率", "增长", "下降", "分析", "统计", "汇总", "明细"]
            for word in common_words:
                if word in placeholder_name:
                    keywords.append(word)

        return keywords

    def _classify_placeholder_type(self, placeholder_name: str) -> str:
        """
        根据占位符名称分类类型

        Args:
            placeholder_name: 占位符名称

        Returns:
            占位符类型：统计类/周期类/图表类
        """
        name_lower = placeholder_name.lower()

        # 图表类关键词
        chart_keywords = [
            "图", "chart", "graph", "plot", "柱状图", "折线图", "饼图", "散点图", "条形图"
        ]
        if any(keyword in name_lower for keyword in chart_keywords):
            return "图表类"

        # 周期类/日期类关键词
        period_keywords = [
            "日期", "时间", "年", "月", "日", "周", "季度", "period", "date", "time",
            "year", "month", "day", "week", "quarter", "当前", "本月", "今年"
        ]
        if any(keyword in name_lower for keyword in period_keywords):
            return "周期类"

        # 统计类关键词（默认类型，用于需要SQL查询的数据）
        stat_keywords = [
            "金额", "数量", "总计", "合计", "平均", "最大", "最小", "增长", "下降",
            "比率", "百分比", "统计", "分析", "明细", "汇总", "count", "sum",
            "avg", "max", "min", "total", "amount", "quantity"
        ]
        if any(keyword in name_lower for keyword in stat_keywords):
            return "统计类"

        # 默认归类为统计类（需要通过SQL获取数据）
        return "统计类"

    def _infer_semantic_type(self, placeholder_name: str) -> Optional[str]:
        """细分语义类型：ranking/compare/period/chart/stat"""
        name = placeholder_name or ""
        name_lower = name.lower()
        # compare
        if any(k in name for k in ["同比", "环比"]) or any(k in name_lower for k in ["yoy", "mom", "compare"]):
            return "compare"
        # ranking
        if any(k in name for k in ["排名", "排行", "最高", "最低", "榜"]) or "top" in name_lower:
            return "ranking"
        import re
        if re.search(r"(前|后)\s*\d+", name):
            return "ranking"
        # chart
        if any(k in name for k in ["图", "chart", "走势图", "折线图", "柱状图", "饼图", "散点图", "条形图"]):
            return "chart"
        # period（弱判断）
        if any(k in name for k in ["月", "周", "季度", "年", "趋势", "周期"]):
            return "period"
        # stat
        if any(k in name for k in ["金额", "数量", "总计", "合计", "平均", "最大", "最小", "统计", "汇总"]):
            return "stat"
        return None

    def _extract_top_n(self, placeholder_name: str) -> Optional[int]:
        """从占位符名称中提取Top N，如 Top10 / 前10 / 后5 等"""
        name = placeholder_name or ""
        import re
        # Top10 / top 5
        m = re.search(r"top\s*(\d+)", name, flags=re.IGNORECASE)
        if m:
            try:
                return int(m.group(1))
            except Exception:
                return None
        # 前10 / 后5
        m = re.search(r"(前|后)\s*(\d+)", name)
        if m:
            try:
                return int(m.group(2))
            except Exception:
                return None
        return None

    def _parse_datetime(self, date_str: Optional[str]) -> Optional[datetime]:
        """解析日期时间字符串"""
        if not date_str:
            return None

        try:
            # 尝试多种日期格式
            for fmt in ["%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"]:
                try:
                    return datetime.strptime(date_str, fmt)
                except ValueError:
                    continue
            return None
        except Exception:
            return None

    async def get_placeholder_contexts(
        self,
        user_id: str,
        template_id: str
    ) -> Dict[str, Any]:
        """
        获取模板的所有占位符上下文信息

        Args:
            user_id: 用户ID
            template_id: 模板ID

        Returns:
            占位符上下文信息字典
        """
        try:
            context_result = await self.build_template_context(
                user_id=user_id,
                template_id=template_id
            )

            if not context_result.get("success"):
                return context_result

            template_context = context_result["template_context"]

            return {
                "success": True,
                "template_id": template_id,
                "placeholder_contexts": template_context["placeholder_contexts"],
                "summary": template_context["summary"],
                "template_info": template_context["template_info"]
            }

        except Exception as e:
            return {"success": False, "error": f"Failed to get placeholder contexts: {str(e)}"}

    async def get_placeholder_context_by_name(
        self,
        user_id: str,
        template_id: str,
        placeholder_name: str
    ) -> Dict[str, Any]:
        """
        获取特定占位符的上下文信息

        Args:
            user_id: 用户ID
            template_id: 模板ID
            placeholder_name: 占位符名称

        Returns:
            特定占位符的上下文信息
        """
        try:
            contexts_result = await self.get_placeholder_contexts(user_id, template_id)

            if not contexts_result.get("success"):
                return contexts_result

            placeholder_contexts = contexts_result["placeholder_contexts"]

            # 查找指定的占位符
            for context in placeholder_contexts:
                if context["placeholder_name"] == placeholder_name:
                    return {
                        "success": True,
                        "placeholder_context": context
                    }

            return {
                "success": False,
                "error": f"Placeholder '{placeholder_name}' not found in template {template_id}"
            }

        except Exception as e:
            return {"success": False, "error": f"Failed to get placeholder context: {str(e)}"}
